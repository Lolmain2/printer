"""Extracts order lines (barcode + quantity + product name) from an order document."""
from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

import openpyxl

from .config import LABELS_PER_SHEET, ORDER_ID_PATTERN, ORDER_LINE_PATTERN


@dataclass
class OrderLine:
    barcode: str
    name: str
    quantity: float
    position: Optional[str] = None

    @property
    def sheets_needed(self) -> int:
        """Number of full label sheets to print (65 stickers per sheet)."""
        return math.ceil(self.quantity / LABELS_PER_SHEET)


@dataclass
class OrderDocument:
    order_id: Optional[str]
    lines: List[OrderLine]
    source: Path


def _parse_bosnian_number(raw: str) -> float:
    """"35,000" -> 35.0 ; "1.234,50" -> 1234.50"""
    cleaned = raw.strip().replace(".", "").replace(",", ".")
    return float(cleaned)


def parse_order_text(text: str) -> List[OrderLine]:
    lines = []
    for match in ORDER_LINE_PATTERN.finditer(text):
        lines.append(
            OrderLine(
                barcode=match.group("barcode"),
                name=" ".join(match.group("name").split()),
                quantity=_parse_bosnian_number(match.group("qty")),
                position=match.group("pos"),
            )
        )
    return lines


def find_order_id(text: str) -> Optional[str]:
    match = ORDER_ID_PATTERN.search(text)
    return match.group(0) if match else None


def _parse_pdf(path: Path) -> OrderDocument:
    from PyPDF2 import PdfReader

    reader = PdfReader(str(path))
    full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return OrderDocument(
        order_id=find_order_id(full_text),
        lines=parse_order_text(full_text),
        source=path,
    )


_BARCODE_HEADERS = {"barcode", "sifra", "šifra", "ean"}
_QTY_HEADERS = {"kolicina", "količina", "kolièina", "quantity", "kom", "kolicina.1"}
_NAME_HEADERS = {"naziv", "name", "proizvod", "artikal"}


def _normalize_header(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _parse_excel(path: Path) -> OrderDocument:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    header_row_idx = None
    col_map = {}
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(20, ws.max_row)), start=1):
        headers = {_normalize_header(c.value): c.column for c in row}
        barcode_col = next((headers[h] for h in _BARCODE_HEADERS if h in headers), None)
        qty_col = next((headers[h] for h in _QTY_HEADERS if h in headers), None)
        if barcode_col and qty_col:
            header_row_idx = row_idx
            col_map = {
                "barcode": barcode_col,
                "qty": qty_col,
                "name": next((headers[h] for h in _NAME_HEADERS if h in headers), None),
            }
            break

    if header_row_idx is None:
        raise ValueError(
            f"Could not find columns for barcode/quantity in {path.name}. "
            f"Expected a header row containing one of {_BARCODE_HEADERS} and one of {_QTY_HEADERS}."
        )

    lines = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row):
        barcode_cell = row[col_map["barcode"] - 1].value
        qty_cell = row[col_map["qty"] - 1].value
        if barcode_cell is None or qty_cell is None:
            continue
        name_cell = row[col_map["name"] - 1].value if col_map.get("name") else ""
        try:
            qty = float(qty_cell)
        except (TypeError, ValueError):
            continue
        barcode = str(barcode_cell).strip()
        if not barcode:
            continue
        lines.append(OrderLine(barcode=barcode, name=str(name_cell or "").strip(), quantity=qty))

    wb.close()
    order_id = None
    if lines:
        order_id = find_order_id(str(path.stem))
    return OrderDocument(order_id=order_id, lines=lines, source=path)


def parse_order_file(path: Path) -> OrderDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return _parse_excel(path)
    raise ValueError(f"Unsupported order file type: {suffix}")
